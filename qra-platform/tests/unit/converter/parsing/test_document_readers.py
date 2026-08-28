from __future__ import annotations

import hashlib
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from qra_converter.image_processing.preprocess import preprocess_image
from qra_converter.image_processing.quality import assess_image_quality
from qra_converter.ocr.fixture_provider import FixtureOcrProvider
from qra_converter.ocr.ports import (
    OcrCell,
    OcrResponse,
    OcrTable,
    OcrTextBlock,
)
from qra_converter.parsing.contracts import BoundingBox
from qra_converter.parsing.pipeline import ParsingPipeline
from qra_converter.readers import CsvReader, DocxReader, XlsReader, XlsxReader

PROJECT_ROOT = Path(__file__).resolve().parents[4]
PARSING_FIXTURES = PROJECT_ROOT / "tests" / "fixtures" / "parsing"


def write_native_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td (Pipeline inspection report with native searchable text) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


class DeterministicReaderTests(unittest.TestCase):
    def test_exif_orientation_is_traceable_and_transform_is_reversible(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "rotated.jpg"
            image = Image.new("RGB", (100, 50), "white")
            exif = Image.Exif()
            exif[274] = 6
            image.save(path, exif=exif)
            processed = preprocess_image(path.read_bytes())
        self.assertEqual(processed.exif_orientation, 6)
        self.assertEqual((processed.width, processed.height), (50, 100))
        original_point = (25.0, 10.0)
        processed_point = processed.original_to_processed.point(*original_point)
        restored = processed.original_to_processed.inverse().point(*processed_point)
        self.assertAlmostEqual(restored[0], original_point[0])
        self.assertAlmostEqual(restored[1], original_point[1])

    def test_image_quality_detects_skew_and_border_clipping(self) -> None:
        page = Image.new("L", (800, 600), "white")
        drawing = ImageDraw.Draw(page)
        for index in range(6):
            drawing.rectangle((100, 100 + index * 60, 700, 110 + index * 60), fill="black")
        rotated = page.rotate(
            5,
            resample=Image.Resampling.BICUBIC,
            fillcolor="white",
        )
        self.assertTrue(assess_image_quality(rotated).skewed)

        clipped = Image.new("L", (800, 600), "white")
        ImageDraw.Draw(clipped).rectangle((0, 100, 20, 500), fill="black")
        self.assertTrue(assess_image_quality(clipped).clipped_content)

    def test_image_quality_detects_blur_exposure_and_low_resolution(self) -> None:
        self.assertTrue(assess_image_quality(Image.new("L", (2000, 2000), 128)).blurry)
        self.assertTrue(assess_image_quality(Image.new("L", (800, 600), 255)).overexposed)
        self.assertTrue(assess_image_quality(Image.new("L", (800, 600), 0)).underexposed)
        self.assertTrue(assess_image_quality(Image.new("L", (599, 800), 128)).low_resolution)

    def test_pipeline_cache_rebinds_artifact_and_preserves_parse_hash(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "cache.csv"
            path.write_text("编号,值\n001,0\n", encoding="utf-8")
            cache = root / "cache"
            first = ParsingPipeline(output_root=root / "first", cache_root=cache).parse_path(path)
            second = ParsingPipeline(output_root=root / "second", cache_root=cache).parse_path(path)
        self.assertFalse(first.cache_hit)
        self.assertTrue(second.cache_hit)
        self.assertEqual(first.document.parse_sha256, second.document.parse_sha256)
        self.assertEqual(second.document.tables[0].cells[2].raw_value, "001")
        provenance = second.document.metadata["parsing_provenance"]
        self.assertEqual(provenance["source_sha256"], second.document.source.checksum_sha256)
        self.assertEqual(provenance["parser_id"], "csv/stdlib")
        self.assertEqual(provenance["parser_version"], second.document.parser_version)
        self.assertEqual(provenance["contract_version"], second.document.contract_version)
        self.assertEqual(len(provenance["cache_key"]), 64)
        self.assertEqual(second.preview_manifest["parsing_provenance"], provenance)

    def test_csv_preserves_encoding_physical_lines_leading_zero_and_empty_string(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "evidence.csv"
            path.write_bytes('编号;说明;空值\r\n001;"第一行\r\n第二行";\r\n'.encode("gb18030"))
            document = (
                CsvReader()
                .parse(
                    path,
                    __import__(
                        "qra_converter.parsing.registry", fromlist=["default_context"]
                    ).default_context(path, CsvReader(), "text/csv"),
                )
                .document
            )
        table = document.tables[0]
        self.assertEqual(document.metadata["encoding"], "gb18030")
        self.assertEqual(table.metadata["record_physical_line_ranges"][1], [2, 3])
        values = {(cell.row_index, cell.column_index): cell for cell in table.cells}
        self.assertEqual(values[(1, 0)].raw_value, "001")
        self.assertEqual(values[(1, 2)].value_type, "EMPTY_STRING")

    def test_xlsx_preserves_formula_merge_hidden_and_date_evidence(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "book.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "管段台账"
            sheet["A1"] = "合并标题"
            sheet.merge_cells("A1:B1")
            sheet["A2"] = "001"
            sheet["B2"] = "=1+1"
            sheet["C2"] = date(2026, 8, 26)
            sheet["A3"] = "隐藏"
            sheet.row_dimensions[3].hidden = True
            sheet.column_dimensions["C"].hidden = True
            sheet.freeze_panes = "B2"
            workbook.save(path)
            workbook.close()
            document = (
                XlsxReader()
                .parse(
                    path,
                    __import__(
                        "qra_converter.parsing.registry", fromlist=["default_context"]
                    ).default_context(
                        path,
                        XlsxReader(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                )
                .document
            )
            legacy_table = XlsxReader().read(path)[0]
        table = document.tables[0]
        cells = {cell.address: cell for cell in table.cells}
        self.assertEqual(cells["A1"].column_span, 2)
        self.assertNotIn("B1", cells)
        self.assertEqual(cells["B2"].formula_text, "=1+1")
        self.assertIsNone(cells["B2"].cached_value)
        self.assertEqual(cells["C2"].parsed_value.date(), date(2026, 8, 26))
        self.assertIn(3, table.metadata["hidden_rows"])
        self.assertIn("C", table.metadata["hidden_columns"])
        codes = {issue.code for issue in document.issues}
        self.assertIn("PARSE.MERGED_CELL_EXPANDED", codes)
        self.assertIn("PARSE.FORMULA_VALUE_MISSING", codes)
        self.assertEqual(legacy_table.metadata["merged_expansions"][0]["master_address"], "A1")

    def test_xls_golden_fixture_preserves_sheets_formula_merge_date_and_leading_zero(self) -> None:
        path = PARSING_FIXTURES / "stage3_legacy.xls"
        reader = XlsReader()
        document = reader.parse(
            path,
            __import__(
                "qra_converter.parsing.registry", fromlist=["default_context"]
            ).default_context(path, reader, "application/vnd.ms-excel"),
        ).document
        cells = {cell.address: cell for cell in document.tables[0].cells}
        self.assertEqual(document.document_kind, "XLS")
        self.assertEqual(document.page_count, 2)
        self.assertEqual(cells["R1C1"].column_span, 4)
        self.assertEqual(cells["R3C1"].raw_value, "001")
        self.assertEqual(cells["R3C2"].formula_text, "=1.0+1.0")
        self.assertEqual(cells["R3C2"].cached_value, 2.0)
        self.assertEqual(cells["R3C3"].parsed_value.date(), date(2026, 8, 28))

    def test_docx_uses_structure_locations_and_never_fabricates_page_numbers(self) -> None:
        xml = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>完整性评价</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>编号</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>"""
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "report.docx"
            with ZipFile(path, "w") as package:
                package.writestr("word/document.xml", xml)
                package.writestr("word/comments.xml", "<comments/>")
            reader = DocxReader()
            document = reader.parse(
                path,
                __import__(
                    "qra_converter.parsing.registry", fromlist=["default_context"]
                ).default_context(path, reader, reader.media_types.__iter__().__next__()),
            ).document
        self.assertEqual(document.page_count, 0)
        self.assertIsNone(document.text_blocks[0].page_number)
        self.assertIn("word/document.xml;paragraph:1", document.text_blocks[0].structure_location)
        self.assertEqual(document.text_blocks[0].style_hint["heading_level"], 1)
        self.assertIn(
            "PARSE.UNSUPPORTED_EMBEDDED_OBJECT",
            {issue.code for issue in document.issues},
        )

    def test_image_pipeline_uses_fixture_ocr_and_flags_low_confidence(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "scan.png"
            image = Image.new("RGB", (800, 600), "white")
            ImageDraw.Draw(image).text((50, 50), "QRA 001", fill="black")
            image.save(path)
            source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            response = OcrResponse(
                "fixture",
                "fixture-v1",
                text_blocks=(
                    OcrTextBlock("管段001", BoundingBox(50, 40, 120, 30), 0.6, "zh-Hans"),
                ),
                tables=(
                    OcrTable(
                        (
                            OcrCell(0, 0, "编号", BoundingBox(50, 100, 100, 40), 0.9),
                            OcrCell(1, 0, "001", BoundingBox(50, 140, 100, 40), 0.9),
                        ),
                        2,
                        1,
                        BoundingBox(50, 100, 100, 80),
                        0.9,
                    ),
                ),
                raw_response_sha256="d" * 64,
                provider_request_id="FIXTURE-1",
            )
            execution = ParsingPipeline(
                output_root=root / "out",
                ocr_provider=FixtureOcrProvider({f"{source_hash}:page-1": response}),
            ).parse_path(path)
        self.assertTrue(execution.succeeded)
        self.assertEqual(execution.document.pages[0].text_blocks[0].text, "管段001")
        self.assertEqual(execution.document.tables[0].cells[1].raw_value, "001")
        self.assertIn(
            "PARSE.LOW_TEXT_CONFIDENCE",
            {issue.code for issue in execution.document.issues},
        )
        self.assertTrue(execution.preview_manifest["resources"])

    def test_image_without_provider_fails_with_explicit_codes_and_no_fake_text(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "scan.jpg"
            Image.new("RGB", (800, 600), "white").save(path)
            execution = ParsingPipeline(output_root=root / "out").parse_path(path)
        self.assertFalse(execution.succeeded)
        self.assertFalse(execution.document.pages[0].text_blocks)
        codes = {issue.code for issue in execution.document.issues}
        self.assertIn("PARSE.OCR_REQUIRED", codes)
        self.assertIn("PARSE.OCR_PROVIDER_NOT_CONFIGURED", codes)

    def test_pdf_routes_native_text_and_scan_pages_by_page(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            native_path = root / "native.pdf"
            write_native_pdf(native_path)
            native = ParsingPipeline(output_root=root / "native-out").parse_path(native_path)
            self.assertTrue(native.succeeded)
            self.assertEqual(native.document.pages[0].classification, "TEXT_NATIVE")
            self.assertEqual(native.document.pages[0].page_number, 1)
            self.assertTrue(native.document.pages[0].text_blocks[0].bbox)

            scan_path = root / "scan.pdf"
            scan_image = Image.new("RGB", (800, 1000), "white")
            ImageDraw.Draw(scan_image).text((80, 100), "SCANNED QRA EVIDENCE", fill="black")
            scan_image.save(scan_path, "PDF", resolution=150)
            source_hash = hashlib.sha256(scan_path.read_bytes()).hexdigest()
            response = OcrResponse(
                "fixture",
                "fixture-v1",
                text_blocks=(OcrTextBlock("扫描证据", BoundingBox(80, 100, 200, 40), 0.95),),
                raw_response_sha256="e" * 64,
            )
            scan = ParsingPipeline(
                output_root=root / "scan-out",
                ocr_provider=FixtureOcrProvider({f"{source_hash}:page-1:image-1": response}),
            ).parse_path(scan_path)
        self.assertTrue(scan.succeeded)
        self.assertEqual(scan.document.pages[0].classification, "SCAN")
        self.assertEqual(scan.document.pages[0].text_blocks[0].text, "扫描证据")
        self.assertEqual(scan.document.pages[0].text_blocks[0].page_number, 1)
        self.assertTrue(scan.document.pages[0].text_blocks[0].bbox)

    def test_mixed_pdf_routes_native_and_scan_pages_in_one_document(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            native_path = root / "native.pdf"
            scan_path = root / "scan.pdf"
            mixed_path = root / "mixed.pdf"
            write_native_pdf(native_path)
            scan_image = Image.new("RGB", (800, 1000), "white")
            ImageDraw.Draw(scan_image).text((80, 100), "MIXED SCAN PAGE", fill="black")
            scan_image.save(scan_path, "PDF", resolution=150)
            writer = PdfWriter()
            writer.append(PdfReader(native_path))
            writer.append(PdfReader(scan_path))
            with mixed_path.open("wb") as output:
                writer.write(output)
            source_hash = hashlib.sha256(mixed_path.read_bytes()).hexdigest()
            response = OcrResponse(
                "fixture",
                "fixture-v1",
                text_blocks=(OcrTextBlock("混合PDF扫描页", BoundingBox(80, 100, 240, 40), 0.96),),
                raw_response_sha256="f" * 64,
                provider_request_id="FIXTURE-MIXED-1",
            )
            execution = ParsingPipeline(
                output_root=root / "out",
                ocr_provider=FixtureOcrProvider(
                    {f"{source_hash}:page-2:image-1": response}
                ),
            ).parse_path(mixed_path)
        self.assertTrue(execution.succeeded)
        self.assertEqual(
            [page.classification for page in execution.document.pages],
            ["TEXT_NATIVE", "SCAN"],
        )
        self.assertEqual(execution.document.pages[1].text_blocks[0].page_number, 2)
        self.assertEqual(
            execution.document.metadata["ocr_calls"][0]["provider_request_id"],
            "FIXTURE-MIXED-1",
        )

    def test_rotated_scan_pdf_preserves_rotation_and_valid_coordinates(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            original_path = root / "scan.pdf"
            rotated_path = root / "rotated.pdf"
            scan_image = Image.new("RGB", (800, 1000), "white")
            ImageDraw.Draw(scan_image).text((100, 120), "ROTATED SCAN", fill="black")
            scan_image.save(original_path, "PDF", resolution=150)
            writer = PdfWriter()
            page = PdfReader(original_path).pages[0]
            page.rotate(90)
            writer.add_page(page)
            with rotated_path.open("wb") as output:
                writer.write(output)
            source_hash = hashlib.sha256(rotated_path.read_bytes()).hexdigest()
            response = OcrResponse(
                "fixture",
                "fixture-v1",
                text_blocks=(OcrTextBlock("旋转扫描页", BoundingBox(100, 120, 200, 40), 0.97),),
                raw_response_sha256="a" * 64,
                provider_request_id="FIXTURE-ROTATED-1",
            )
            execution = ParsingPipeline(
                output_root=root / "out",
                ocr_provider=FixtureOcrProvider(
                    {f"{source_hash}:page-1:image-1": response}
                ),
            ).parse_path(rotated_path)
        block = execution.document.pages[0].text_blocks[0]
        self.assertTrue(execution.succeeded)
        self.assertEqual(execution.document.pages[0].metadata["rotation"], 90)
        self.assertGreater(block.bbox.width, 0)
        self.assertGreater(block.bbox.height, 0)
        self.assertLessEqual(block.bbox.right, block.page_width)
        self.assertLessEqual(block.bbox.bottom, block.page_height)

    def test_corrupt_empty_and_oversized_inputs_return_structured_errors(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            corrupt = root / "corrupt.pdf"
            corrupt.write_bytes(b"%PDF-1.7\ncorrupt")
            corrupt_execution = ParsingPipeline(output_root=root / "corrupt-out").parse_path(
                corrupt
            )

            empty = root / "empty.csv"
            empty.write_bytes(b"")
            empty_execution = ParsingPipeline(output_root=root / "empty-out").parse_path(empty)

            oversized = root / "oversized.png"
            Image.new("RGB", (20, 20), "white").save(oversized)
            with patch("qra_converter.image_processing.preprocess.MAX_IMAGE_PIXELS", 100):
                oversized_execution = ParsingPipeline(
                    output_root=root / "oversized-out"
                ).parse_path(oversized)

        self.assertIn(
            "PARSE.DOCUMENT_CORRUPT",
            {issue.code for issue in corrupt_execution.document.issues},
        )
        self.assertIn(
            "PARSE.NO_CONTENT",
            {issue.code for issue in empty_execution.document.issues},
        )
        self.assertIn(
            "PARSE.DOCUMENT_CORRUPT",
            {issue.code for issue in oversized_execution.document.issues},
        )

    def test_encrypted_pdf_is_rejected_without_password_attempt(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "encrypted.pdf"
            writer = PdfWriter()
            writer.add_blank_page(200, 200)
            writer.encrypt("secret")
            with path.open("wb") as output:
                writer.write(output)
            execution = ParsingPipeline(output_root=root / "out").parse_path(path)
        self.assertFalse(execution.succeeded)
        self.assertIn(
            "PARSE.ENCRYPTED_UNSUPPORTED",
            {issue.code for issue in execution.document.issues},
        )


if __name__ == "__main__":
    unittest.main()
